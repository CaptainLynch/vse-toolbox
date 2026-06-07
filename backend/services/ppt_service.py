import io
import json
import logging
import re
import sys
import textwrap
from datetime import datetime, timedelta
from pathlib import Path

logger = logging.getLogger("VSE_TOOLBOX.ppt")

from config import BASE_DIR, DATA_ROOT

TEMPLATES_DIR = BASE_DIR / "templates"
MASTER_DIR = TEMPLATES_DIR / "master"
CONFIG_DIR = TEMPLATES_DIR / "config"

OUTPUT_DIR = DATA_ROOT / "templates" / "output"

# ---------------------------------------------------------------------------
# 公司配色常量
# ---------------------------------------------------------------------------

COLORS = {
    "primary": "#d4af37",
    "bg": "#0a0a0c",
    "bg_light": "#1a1a1f",
    "text": "#ffffff",
    "text_secondary": "#8a8f98",
    "danger": "#ff4d4d",
    "success": "#4ade80",
    "primary_rgb": (212 / 255, 175 / 255, 55 / 255),
    "bg_rgb": (10 / 255, 10 / 255, 12 / 255),
    "danger_rgb": (1.0, 77 / 255, 77 / 255),
    "success_rgb": (74 / 255, 222 / 255, 128 / 255),
}

# ---------------------------------------------------------------------------
# matplotlib 后端配置（打包后也必须是非交互式）
# ---------------------------------------------------------------------------

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
from matplotlib.patches import FancyBboxPatch
import numpy as np


# ---------------------------------------------------------------------------
# 4.1 ChartGenerator — 暗色主题图表生成
# ---------------------------------------------------------------------------

class ChartGenerator:
    """生成暗色主题 PNG 图表，用于插入 PPT。"""

    @staticmethod
    def _setup_dark_style():
        plt.rcParams.update({
            "figure.facecolor": COLORS["bg"],
            "axes.facecolor": COLORS["bg_light"],
            "axes.edgecolor": COLORS["text_secondary"],
            "axes.labelcolor": COLORS["text"],
            "text.color": COLORS["text"],
            "xtick.color": COLORS["text_secondary"],
            "ytick.color": COLORS["text_secondary"],
            "grid.color": "#2a2a30",
            "grid.alpha": 0.6,
            "legend.facecolor": COLORS["bg_light"],
            "legend.edgecolor": COLORS["text_secondary"],
            "legend.labelcolor": COLORS["text_secondary"],
            "font.sans-serif": ["Microsoft YaHei", "SimHei", "DejaVu Sans"],
            "axes.unicode_minus": False,
        })

    @staticmethod
    def issue_trend(dates: list[str], counts: list[int], output_path: Path) -> Path:
        """问题趋势折线图。"""
        ChartGenerator._setup_dark_style()
        fig, ax = plt.subplots(figsize=(10, 4.5))
        x = range(len(dates))
        ax.plot(x, counts, color=COLORS["primary"], linewidth=2.5, marker="o",
                markersize=6, markerfacecolor=COLORS["primary"])
        ax.fill_between(x, counts, alpha=0.15, color=COLORS["primary"])
        ax.set_xticks(x)
        ax.set_xticklabels(dates, rotation=45, ha="right", fontsize=8)
        ax.set_ylabel("Issue Count", fontsize=10, color=COLORS["text_secondary"])
        ax.yaxis.set_major_locator(mticker.MaxNLocator(integer=True))
        ax.grid(axis="y", linestyle="--")
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        fig.tight_layout()
        fig.savefig(output_path, dpi=150, facecolor=COLORS["bg"], edgecolor="none")
        plt.close(fig)
        return output_path

    @staticmethod
    def department_closed_rate(departments: list[str], rates: list[int],
                               output_path: Path) -> Path:
        """部门关闭率横向柱状图。"""
        ChartGenerator._setup_dark_style()
        fig, ax = plt.subplots(figsize=(9, len(departments) * 0.7 + 1.5))
        y_pos = range(len(departments))
        colors_bar = [COLORS["success"] if r >= 60 else
                      COLORS["primary"] if r >= 40 else
                      COLORS["danger"] for r in rates]
        ax.barh(y_pos, rates, height=0.5, color=colors_bar, edgecolor="none")
        ax.set_yticks(y_pos)
        ax.set_yticklabels(departments, fontsize=9)
        ax.set_xlabel("Close Rate %", fontsize=10, color=COLORS["text_secondary"])
        ax.set_xlim(0, 105)
        for i, r in enumerate(rates):
            ax.text(r + 1, i, f"{r}%", va="center", fontsize=9,
                    color=COLORS["text"])
        ax.invert_yaxis()
        ax.grid(axis="x", linestyle="--", alpha=0.4)
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        fig.tight_layout()
        fig.savefig(output_path, dpi=150, facecolor=COLORS["bg"], edgecolor="none")
        plt.close(fig)
        return output_path

    @staticmethod
    def priority_pie(p0: int, p1: int, p2: int, p3: int, output_path: Path) -> Path:
        """优先级分布饼图。"""
        ChartGenerator._setup_dark_style()
        labels = ["P0 - Critical", "P1 - High", "P2 - Medium", "P3 - Low"]
        sizes = [p0, p1, p2, p3]
        pie_colors = [COLORS["danger"], "#ff8c42", COLORS["primary"], COLORS["text_secondary"]]
        explode = (0.05, 0.02, 0, 0)

        fig, ax = plt.subplots(figsize=(7, 5))
        if sum(sizes) == 0:
            ax.text(0.5, 0.5, "No Issues", ha="center", va="center",
                    fontsize=20, color=COLORS["text_secondary"], transform=ax.transAxes)
            ax.axis("off")
        else:
            wedges, texts, autotexts = ax.pie(
                sizes, explode=explode, labels=None, colors=pie_colors,
                autopct="%1.1f%%", startangle=140, pctdistance=0.75,
                wedgeprops={"edgecolor": COLORS["bg"], "linewidth": 2},
            )
            for at in autotexts:
                at.set_fontsize(9)
                at.set_color(COLORS["bg"])
            ax.legend(wedges, labels, loc="lower right", fontsize=8,
                      bbox_to_anchor=(1.35, 0.5))
        fig.tight_layout()
        fig.savefig(output_path, dpi=150, facecolor=COLORS["bg"], edgecolor="none")
        plt.close(fig)
        return output_path

    @staticmethod
    def weekly_overview(metrics: dict, output_path: Path) -> Path:
        """周报概览指标卡（4 个 KPI 卡片并排）。"""
        ChartGenerator._setup_dark_style()
        fig, axes = plt.subplots(1, 4, figsize=(12, 3))
        items = [
            ("Open Issues", metrics.get("total_open", 0), COLORS["text"]),
            ("New This Week", metrics.get("new_this_week", 0), COLORS["primary"]),
            ("Closed This Week", metrics.get("closed_this_week", 0), COLORS["success"]),
            ("High Risk (P0)", metrics.get("high_risk_count", 0), COLORS["danger"]),
        ]
        for ax, (label, value, color) in zip(axes, items):
            ax.set_facecolor(COLORS["bg_light"])
            ax.text(0.5, 0.55, str(value), transform=ax.transAxes,
                    fontsize=36, fontweight="bold", color=color,
                    ha="center", va="center")
            ax.text(0.5, 0.15, label, transform=ax.transAxes,
                    fontsize=10, color=COLORS["text_secondary"],
                    ha="center", va="center")
            ax.set_xticks([])
            ax.set_yticks([])
            for spine in ax.spines.values():
                spine.set_color(COLORS["text_secondary"])
                spine.set_linewidth(0.5)
        fig.tight_layout()
        fig.savefig(output_path, dpi=150, facecolor=COLORS["bg"], edgecolor="none")
        plt.close(fig)
        return output_path


# ---------------------------------------------------------------------------
# 4.2 TemplateEngine — PPT 母版填充引擎
# ---------------------------------------------------------------------------

PLACEHOLDER_RE = re.compile(r"\{\{(\w+)\}\}")


class TemplateEngine:
    """加载母版 .pptx，解析 {{placeholder}}，执行文本/表格/图表替换。"""

    def __init__(self, template_name: str, config_name: str):
        from pptx import Presentation
        from pptx.util import Inches, Pt

        self.template_path = MASTER_DIR / template_name
        self.config_path = CONFIG_DIR / config_name
        self.config: dict = {}
        self.prs: "Presentation" | None = None
        self._load_config()

    def _load_config(self):
        if self.config_path.exists():
            with open(self.config_path, "r", encoding="utf-8") as f:
                self.config = json.load(f)
        else:
            self.config = {}

    def load(self):
        from pptx import Presentation
        if not self.template_path.exists():
            raise FileNotFoundError(f"母版模板不存在: {self.template_path}")
        self.prs = Presentation(str(self.template_path))
        return self

    def render(self, data: dict, output_path: Path) -> Path:
        """遍历所有 Slide 和 Shape，替换占位符。"""
        if self.prs is None:
            raise RuntimeError("Template not loaded. Call .load() first.")

        for slide in self.prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    self._replace_text(shape, data)
                elif shape.has_table:
                    self._fill_table(shape, data)

        # 插入图表图片（通过 data["_charts"] 指定位置和图片路径）
        charts = data.get("_charts", [])
        for chart_spec in charts:
            self._insert_chart_image(chart_spec)

        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(output_path))
        logger.info("PPT saved to %s", output_path)
        return output_path

    def _replace_text(self, shape, data: dict):
        """替换 Shape 文本中的 {{placeholder}}。"""
        from pptx.util import Pt
        from pptx.dml.color import RGBColor

        original_text = shape.text
        matches = PLACEHOLDER_RE.findall(original_text)
        if not matches:
            return

        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                replaced = run.text
                for key in matches:
                    value = data.get(key)
                    if value is not None:
                        replaced = replaced.replace(f"{{{{{key}}}}}", str(value))
                if replaced != run.text:
                    run.text = replaced
                    # 应用样式配置
                    for key in matches:
                        if key in self.config and key in data:
                            style = self.config[key]
                            if "font_size" in style:
                                run.font.size = Pt(style["font_size"])
                            if "color" in style:
                                run.font.color.rgb = RGBColor.from_string(style["color"].lstrip("#"))
                            if "bold" in style and style["bold"]:
                                run.font.bold = True

    def _fill_table(self, shape, data: dict):
        """用 data dict 中的 table_xxx 数据填充表格。"""
        from pptx.util import Pt
        from pptx.dml.color import RGBColor

        table = shape.table
        table_data = None

        for para in shape.text_frame.paragraphs if shape.has_text_frame else []:
            for run in para.runs:
                for key in PLACEHOLDER_RE.findall(run.text):
                    if key.startswith("table_") and key in data:
                        table_data = data[key]
                        break

        if table_data is None:
            return

        rows_data = table_data.get("rows", [])
        headers = table_data.get("headers", [])
        style_cfg = self.config.get("table_default", {})

        # 写入表头
        if headers:
            for ci, h in enumerate(headers):
                if ci < len(table.columns):
                    cell = table.cell(0, ci)
                    cell.text = str(h)
                    self._style_cell(cell, bold=True, font_size=style_cfg.get("header_size", 10))

        # 写入数据行
        for ri, row in enumerate(rows_data):
            tbl_row = ri + (1 if headers else 0)
            if tbl_row >= len(table.rows):
                break
            for ci, val in enumerate(row):
                if ci < len(table.columns):
                    cell = table.cell(tbl_row, ci)
                    cell.text = str(val)
                    self._style_cell(cell, font_size=style_cfg.get("body_size", 9))

    @staticmethod
    def _style_cell(cell, bold: bool = False, font_size: int = 9):
        from pptx.util import Pt
        from pptx.dml.color import RGBColor

        for para in cell.text_frame.paragraphs:
            para.font.size = Pt(font_size)
            para.font.color.rgb = RGBColor.from_string(COLORS["text"].lstrip("#"))
            if bold:
                para.font.bold = True

    def _insert_chart_image(self, spec: dict):
        """在指定 slide 和位置插入图表 PNG 图片。"""
        from pptx.util import Inches

        slide_idx = spec.get("slide", 0)
        img_path = spec.get("path")
        left = Inches(spec.get("left", 1))
        top = Inches(spec.get("top", 2.5))
        width = Inches(spec.get("width", 8))
        height = Inches(spec.get("height", 4))

        if not img_path or not Path(img_path).exists():
            logger.warning("Chart image not found: %s", img_path)
            return

        if slide_idx < len(self.prs.slides):
            slide = self.prs.slides[slide_idx]
            slide.shapes.add_picture(str(img_path), left, top, width, height)


# ---------------------------------------------------------------------------
# 4.3 DataAdapter — 数据查询与格式化
# ---------------------------------------------------------------------------

class DataAdapter:
    """从 SQLite 查询数据，组装 render() 所需的 data dict。"""

    @staticmethod
    def weekly_data(week_start: str | None = None,
                    project_name: str | None = None,
                    author: str | None = None) -> dict:
        from services.db import DBManager

        # 确定周范围
        today = datetime.now()
        if week_start:
            ws = datetime.strptime(week_start, "%Y-%m-%d")
        else:
            ws = today.replace(hour=0, minute=0, second=0, microsecond=0)
            while ws.weekday() != 0:
                ws -= timedelta(days=1)
        we = ws + timedelta(days=6)
        week_label = f"{ws.strftime('%m/%d')} - {we.strftime('%m/%d')}"

        # 统计数据
        stats = DBManager.get_issue_stats()
        issues_result, _ = DBManager.list_issues(page=1, size=50, status="open")
        milestones = DBManager.list_milestones()

        # 图表数据
        chart_dir = OUTPUT_DIR / "_charts"
        chart_dir.mkdir(parents=True, exist_ok=True)

        # 趋势图
        trend_dates = [t["date"] for t in stats["trend"]]
        trend_counts = [t["count"] for t in stats["trend"]]
        trend_path = chart_dir / f"trend_{today.strftime('%Y%m%d_%H%M%S')}.png"
        ChartGenerator.issue_trend(trend_dates, trend_counts, trend_path)

        # 部门关闭率图
        dept_names = [d["department"] for d in stats["department_stats"]]
        dept_rates = [d["closed_rate"] for d in stats["department_stats"]]
        dept_path = chart_dir / f"dept_{today.strftime('%Y%m%d_%H%M%S')}.png"
        ChartGenerator.department_closed_rate(dept_names, dept_rates, dept_path)

        # 概览 KPI 图
        overview_path = chart_dir / f"overview_{today.strftime('%Y%m%d_%H%M%S')}.png"
        ChartGenerator.weekly_overview(stats, overview_path)

        # 问题表格
        table_headers = ["ID", "Priority", "Component", "Department", "Assignee", "Status"]
        table_rows = []
        for item in issues_result:
            table_rows.append([
                item.get("id", ""),
                item.get("priority", ""),
                item.get("component", ""),
                item.get("department", ""),
                item.get("assignee", "") or "-",
                item.get("status", ""),
            ])

        # 里程碑表格
        ms_headers = ["Milestone", "Category", "Progress", "Target Date"]
        ms_rows = []
        for m in milestones:
            ms_rows.append([
                m.get("name", ""),
                m.get("category", ""),
                f"{m.get('percentage', 0)}%",
                m.get("target_date", "-") or "-",
            ])

        data = {
            "text_report_title": f"{project_name or 'PM'} 项目周报",
            "text_week_range": week_label,
            "text_author": author or "项目管理部",
            "text_open_count": str(stats["total_open"]),
            "text_new_count": str(stats["new_this_week"]),
            "text_closed_count": str(stats["closed_this_week"]),
            "text_p0_count": str(stats["high_risk_count"]),
            "table_issues": {"headers": table_headers, "rows": table_rows},
            "table_milestones": {"headers": ms_headers, "rows": ms_rows},
            "_charts": [
                {"slide": 1, "path": str(trend_path), "left": 1, "top": 2.0, "width": 8, "height": 4},
                {"slide": 2, "path": str(overview_path), "left": 0.5, "top": 1.5, "width": 9, "height": 2.5},
                {"slide": 2, "path": str(dept_path), "left": 0.5, "top": 4.2, "width": 9, "height": 3},
            ],
        }
        return data

    @staticmethod
    def deliverable_data(deliverable_name: str, responsible: str | None = None) -> dict:
        from services.db import DBManager

        today = datetime.now()
        stats = DBManager.get_issue_stats()

        # 拉取所有 issues，按 deliverable_name 模糊匹配 component/description
        all_issues, _ = DBManager.list_issues(page=1, size=200)
        issues_result: list[dict] = [
            i for i in all_issues
            if deliverable_name.lower() in i.get("component", "").lower()
            or deliverable_name.lower() in i.get("description", "").lower()
        ]

        # 图表
        chart_dir = OUTPUT_DIR / "_charts"
        chart_dir.mkdir(parents=True, exist_ok=True)

        # 优先级饼图
        p0 = sum(1 for i in issues_result if i.get("priority") == "P0")
        p1 = sum(1 for i in issues_result if i.get("priority") == "P1")
        p2 = sum(1 for i in issues_result if i.get("priority") == "P2")
        p3 = sum(1 for i in issues_result if i.get("priority") == "P3")
        pie_path = chart_dir / f"pie_{today.strftime('%Y%m%d_%H%M%S')}.png"
        ChartGenerator.priority_pie(p0, p1, p2, p3, pie_path)

        # 问题表格
        table_headers = ["ID", "Priority", "Description", "Status", "Assignee"]
        table_rows = []
        for item in issues_result:
            table_rows.append([
                item.get("id", ""),
                item.get("priority", ""),
                item.get("description", ""),
                item.get("status", ""),
                item.get("assignee", "") or "-",
            ])

        data = {
            "text_deliverable_name": deliverable_name,
            "text_responsible": responsible or "-",
            "text_date": today.strftime("%Y-%m-%d"),
            "text_issue_count": str(len(issues_result)),
            "text_p0_count": str(p0),
            "table_issues": {"headers": table_headers, "rows": table_rows},
            "_charts": [
                {"slide": 2, "path": str(pie_path), "left": 1.5, "top": 2.0, "width": 6, "height": 4.5},
            ],
        }
        return data


# ---------------------------------------------------------------------------
# 工厂函数：生成 PPT
# ---------------------------------------------------------------------------

def generate_weekly_ppt(week_start: str | None = None,
                        project_name: str | None = None,
                        author: str | None = None) -> dict:
    """生成周报 PPT 的完整流程。"""
    now = datetime.now()
    data = DataAdapter.weekly_data(week_start, project_name, author)

    engine = TemplateEngine("weekly_report_master.pptx", "weekly_map.json")
    engine.load()

    output_name = f"weekly_report_{now.strftime('%Y%m%d_%H%M%S')}.pptx"
    output_path = OUTPUT_DIR / output_name
    engine.render(data, output_path)

    week_start_dt = datetime.strptime(week_start, "%Y-%m-%d") if week_start else now
    week_end = week_start_dt + timedelta(days=6) if week_start else now + timedelta(days=6 - now.weekday())
    week_date = f"{week_start_dt.strftime('%m/%d')}-{week_end.strftime('%m/%d')}"

    return {
        "file_path": str(output_path),
        "file_name": output_name,
        "week_date": week_date,
    }


def generate_deliverable_ppt(deliverable_name: str,
                             responsible: str | None = None) -> dict:
    """生成交付物报告 PPT 的完整流程。"""
    now = datetime.now()
    data = DataAdapter.deliverable_data(deliverable_name, responsible)

    engine = TemplateEngine("deliverable_master.pptx", "deliverable_map.json")
    engine.load()

    output_name = f"deliverable_{now.strftime('%Y%m%d_%H%M%S')}.pptx"
    output_path = OUTPUT_DIR / output_name
    engine.render(data, output_path)

    return {
        "file_path": str(output_path),
        "file_name": output_name,
    }
